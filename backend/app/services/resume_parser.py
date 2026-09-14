import hashlib
import io
import re

import pypdf
from docx import Document

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_PARSE_ATTEMPTS = 3
MAX_REPARSE_ATTEMPTS = 3

ALLOWED_TYPES = {
    "application/pdf": (b"%PDF", "pdf"),
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": (
        b"PK",
        "docx",
    ),
}


class ResumeParseError(Exception):
    pass


def validate_upload(data: bytes, content_type: str) -> str:
    """Returns the safe file extension or raises ResumeParseError."""
    if not data:
        raise ResumeParseError("File is empty")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ResumeParseError(
            f"File exceeds the {MAX_UPLOAD_BYTES // (1024 * 1024)} MB limit"
        )
    if content_type not in ALLOWED_TYPES:
        raise ResumeParseError("Only PDF and DOCX files are supported")
    magic, extension = ALLOWED_TYPES[content_type]
    if not data.startswith(magic) and content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        raise ResumeParseError("File content does not match its claimed type")
    if content_type == "application/pdf" and not data.startswith(magic):
        raise ResumeParseError("File is not a valid PDF")
    return extension


def extract_text(data: bytes, content_type: str) -> str:
    if content_type == "application/pdf":
        return extract_text_pdf(data)
    return extract_text_docx(data)


def extract_text_pdf(data: bytes) -> str:
    try:
        reader = pypdf.PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ResumeParseError(f"Could not read PDF text: {exc}") from exc
    text = "\n".join(pages).strip()
    if not text:
        raise ResumeParseError("PDF contains no extractable text")
    return text


def extract_text_docx(data: bytes) -> str:
    try:
        document = Document(io.BytesIO(data))
        parts: list[str] = []
        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                parts.append(" | ".join(cell for cell in cells if cell))
    except Exception as exc:
        raise ResumeParseError(f"Could not read DOCX text: {exc}") from exc
    text = "\n".join(parts).strip()
    if not text:
        raise ResumeParseError("DOCX contains no extractable text")
    return text


def bounded_parse_resume(
    data: bytes, content_type: str, attempts: int = MAX_PARSE_ATTEMPTS
) -> tuple[str, dict]:
    """Parse with bounded internal attempts; never retries indefinitely."""
    last_error: str = "Unknown parse failure"
    for _ in range(max(1, attempts)):
        try:
            text = extract_text(data, content_type)
            return text, extract_fields(text)
        except ResumeParseError as exc:
            last_error = str(exc)
    raise ResumeParseError(last_error)


_SKILL_LEXICON = [
    "Python",
    "JavaScript",
    "TypeScript",
    "React",
    "React Native",
    "Node.js",
    "Angular",
    "Vue",
    "Java",
    "C++",
    "C#",
    "Go",
    "Rust",
    "Ruby",
    "PHP",
    "SQL",
    "PostgreSQL",
    "MySQL",
    "MongoDB",
    "Redis",
    "Kafka",
    "GraphQL",
    "REST API",
    "FastAPI",
    "Django",
    "Flask",
    "Spring Boot",
    "Docker",
    "Kubernetes",
    "AWS",
    "GCP",
    "Azure",
    "Terraform",
    "Git",
    "CI/CD",
    "Machine Learning",
    "Deep Learning",
    "TensorFlow",
    "PyTorch",
    "Pandas",
    "NumPy",
    "HTML",
    "CSS",
    "Tailwind",
]

_LANGUAGE_LEXICON = [
    "English",
    "Hindi",
    "Telugu",
    "Tamil",
    "Kannada",
    "Malayalam",
    "Marathi",
    "Bengali",
    "Gujarati",
    "Punjabi",
    "French",
    "German",
    "Spanish",
    "Mandarin Chinese",
    "Japanese",
    "Nepali",
]

_DEGREE_KEYWORDS = [
    "b.tech",
    "b.e.",
    "m.tech",
    "m.sc",
    "m.b.a",
    "mba",
    "ph.d",
    "b.sc",
    "b.com",
    "m.com",
    "bca",
    "mca",
    "bachelor",
    "master",
    "degree",
]


def _lexicon_matches(text: str, lexicon: list[str]) -> list[str]:
    lower = text.lower()
    found: list[str] = []
    positions: list[int] = []
    for entry in lexicon:
        pattern = re.compile(rf"(?<![a-zA-Z0-9]){re.escape(entry.lower())}(?![a-zA-Z0-9])")
        match = pattern.search(lower)
        if match and entry not in found:
            found.append(entry)
            positions.append(match.start())
    return [entry for _, entry in sorted(zip(positions, found))]


_SECTION_HEADINGS = re.compile(
    r"^(experience|professional experience|work experience|employment history|"
    r"education|academic|skills|technical skills|languages|certifications|projects|"
    r"summary|objective|additional|profile|contact)\b",
    re.IGNORECASE,
)


def _section_lines(text: str, wanted: tuple[str, ...]) -> list[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    in_section = False
    section: list[str] = []
    for line in lines:
        match = _SECTION_HEADINGS.match(line)
        if match:
            heading = match.group(1).lower()
            if in_section:
                break
            if any(word in heading for word in wanted):
                in_section = True
                continue
            continue
        if in_section:
            section.append(line)
    return section


def extract_fields(text: str) -> dict:
    email = ""
    match_email = re.search(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}", text)
    if match_email:
        email = match_email.group(0)

    phone = ""
    match_phone = re.search(
        r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{2,5}\)?[\s-]?)?\d{3,5}[\s-]?\d{3,5}",
        text,
    )
    if match_phone:
        candidate = match_phone.group(0).strip()
        digits = re.sub(r"\D", "", candidate)
        if 8 <= len(digits) <= 15:
            phone = candidate

    total_experience = None
    match_exp = re.search(r"(\d{1,2})\+?\s*(?:years|yrs|year)\b", text, re.IGNORECASE)
    if match_exp:
        total_experience = int(match_exp.group(1))

    skills = _lexicon_matches(text, _SKILL_LEXICON)
    languages = _lexicon_matches(text, _LANGUAGE_LEXICON)

    education: list[dict] = []
    for line in _section_lines(text, ("education", "academic")):
        lower = line.lower()
        year = re.search(r"\b(19|20)\d{2}\b", line)
        if any(keyword in lower for keyword in _DEGREE_KEYWORDS) or year:
            education.append(
                {
                    "degree": line,
                    "institution": "",
                    "graduation_year": int(year.group(0)) if year else None,
                }
            )

    experience: list[dict] = []
    for line in _section_lines(text, ("experience", "employment", "professional")):
        for separator in (" at ", " — ", " - "):
            if separator in line:
                job_title, _, employer = line.partition(separator)
                experience.append(
                    {
                        "job_title": job_title.strip(),
                        "employer": employer.strip(),
                        "projects": [],
                    }
                )
                break

    current_role = experience[0]["job_title"] if experience else ""

    return {
        "full_name": "",
        "email": email,
        "phone": phone or None,
        "location": None,
        "current_role": current_role,
        "total_experience": total_experience,
        "languages": languages,
        "certifications": [],
        "work_authorization": None,
        "skills": [{"name": skill, "experience": None} for skill in skills],
        "experience": experience,
        "education": education,
        "preferences": None,
    }


def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()